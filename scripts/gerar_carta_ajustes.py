"""Gera carta_de_ajustes_e_explicacao_das_mudancas.docx (versão concisa, voz dos autores)."""
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Emu

LINK = "[LINK-ANONIMIZADO-PENDENTE]"
TITULO = "Predição de risco de Doença de Chagas Aguda em nível de município-ano com Random Forest e XGBoost"
PARAGRAFOS = [
    "Agradecemos aos avaliadores pelos pareceres, que levaram a mudanças importantes no artigo. Resumimos a seguir o que foi alterado.",
    "A alteração mais relevante diz respeito à validação. A versão submetida apresentava AUC-ROC acima de 0,90 como resultado central, obtida por validação cruzada que mistura anos entre treino e teste. Um dos pareceres apontou que esse desenho poderia superestimar o desempenho em uso real, em que se prevê um ano ainda não observado. Implementamos então uma validação temporal por origem deslizante (treino até o ano t, teste em t+1) e a validação cruzada aninhada. A AUC-ROC cai para cerca de 0,65 sob validação temporal, valor próximo ao de uma regra simples que repete os casos do ano anterior (0,6371), e não melhora ao retirar o atributo ano. A validação aninhada difere da não aninhada em no máximo 0,0007, o que descarta otimismo causado pela busca de hiperparâmetros. O artigo passou a reportar os dois regimes lado a lado, e a discussão e o resumo foram reescritos com uma leitura mais cautelosa: o modelo serve melhor à priorização retrospectiva do que à previsão do ano seguinte.",
    "Corrigimos também erros de construção da base. O código de município do SINAN muda de seis para sete dígitos entre as duas eras do layout, e o pipeline não os unificava, o que duplicava municípios (Pernambuco aparecia com 204, e não 185) e quebrava as defasagens de casos. Todos os códigos foram normalizados e a grade passou a partir da lista oficial de 510 municípios, o que elevou a base de 12.090 para 13.260 município-anos, sem alterar os 259 de risco. A população, que entrava por estado, passou a ser obtida por município na própria tabela 6579 do SIDRA, e a afirmação em contrário foi removida. As Seções 3.1 e 3.2 agora explicitam o campo de município usado (município de notificação, que coincide com o de residência em 93,9% dos registros), o critério de caso confirmado e a deduplicação.",
    "Acrescentamos três referências de comparação (regra do ano anterior e duas regressões logísticas), métricas complementares (AUC-PR, precisão e sensibilidade nos maiores riscos previstos, calibração), mantendo a AUC-ROC como única métrica de decisão, e a importância por permutação, menos sujeita ao viés da importância nativa. A seção de curvas ROC foi ampliada com a leitura da sensibilidade por faixa de falsos positivos e suas limitações. A Seção 3.4 foi condensada e a afirmação sobre o parâmetro scale_pos_weight foi corrigida. A Introdução foi reduzida, e a expressão \"relatório técnico de origem\", que se referia apenas à documentação interna do pipeline e não a publicação anterior, foi removida. Os gráficos, as referências e as citações foram revisados conforme as normas ABNT NBR 6023:2025 e NBR 10520:2023, e o texto foi reformatado segundo o modelo do evento e revisado quanto a repetições de estilo. Conferimos ainda as citações contra as fontes, sem encontrar cópia literal, e reescrevemos as passagens que acompanhavam de perto o material de apoio.",
    f"Atendendo ao pedido de material suplementar, o código do pipeline, os testes e os relatórios estão disponíveis em versão anonimizada em {LINK}, referenciada na Seção 3.6 do artigo. Deixamos para trabalhos futuros a inclusão de atributos socioeconômicos municipais, como proporção de população rural e indicadores de saneamento, por dependerem de dados censitários (2000, 2010 e 2022) que exigem estratégia própria de interpolação para a série anual.",
]
d = docx.Document("carta_de_ajustes_e_explicacao_das_mudancas.docx")
body = d.element.body
for el in list(body):
    if el.tag.endswith("}sectPr"):
        continue
    body.remove(el)

def par(texto, tam, negrito=False, alin=WD_ALIGN_PARAGRAPH.JUSTIFY, depois=8):
    p = d.add_paragraph()
    p.alignment = alin
    p.paragraph_format.space_after = Pt(depois)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(texto)
    r.font.name = "Calibri"
    r.font.size = Pt(tam)
    r.bold = negrito
    return p

par("Carta de ajustes", 15, True, WD_ALIGN_PARAGRAPH.CENTER, 2)
par(TITULO, 12, False, WD_ALIGN_PARAGRAPH.CENTER, 14)
for t in PARAGRAFOS:
    par(t, 11)
d.save("carta_de_ajustes_e_explicacao_das_mudancas.docx")
